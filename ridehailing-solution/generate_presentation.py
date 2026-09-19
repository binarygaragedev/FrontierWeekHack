from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_FILE = BASE_DIR / "ridehailing-multi-agent-presentation-v2-notes.pptx"


TITLE_COLOR = RGBColor(22, 44, 74)
ACCENT_COLOR = RGBColor(0, 120, 170)
TEXT_COLOR = RGBColor(40, 40, 40)
BG_COLOR = RGBColor(243, 247, 252)
CARD_COLOR = RGBColor(255, 255, 255)
MUTED_COLOR = RGBColor(96, 115, 136)


def style_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

    top_band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        Inches(13.333),
        Inches(0.62),
    )
    top_band.fill.solid()
    top_band.fill.fore_color.rgb = RGBColor(13, 34, 59)
    top_band.line.fill.background()


def add_speaker_notes(slide, notes: str) -> None:
    notes_frame = slide.notes_slide.notes_text_frame
    notes_frame.clear()
    notes_frame.text = notes


def add_title_slide(prs: Presentation, title: str, subtitle: str, notes: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    style_background(slide)
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle

    title_frame = slide.shapes.title.text_frame
    title_run = title_frame.paragraphs[0].runs[0]
    title_run.font.size = Pt(42)
    title_run.font.bold = True
    title_run.font.color.rgb = TITLE_COLOR

    sub_frame = slide.placeholders[1].text_frame
    sub_run = sub_frame.paragraphs[0].runs[0]
    sub_run.font.size = Pt(20)
    sub_run.font.color.rgb = ACCENT_COLOR

    add_speaker_notes(slide, notes)


def add_bullet_slide(
    prs: Presentation, title: str, bullets: list[str], notes: str
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    style_background(slide)

    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.55),
        Inches(1.15),
        Inches(12.25),
        Inches(5.95),
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = CARD_COLOR
    panel.line.color.rgb = RGBColor(221, 230, 240)

    slide.shapes.title.text = title

    title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    title_run.font.size = Pt(32)
    title_run.font.bold = True
    title_run.font.color.rgb = TITLE_COLOR

    body = slide.shapes.placeholders[1].text_frame
    body.clear()

    for idx, bullet in enumerate(bullets):
        p = body.paragraphs[0] if idx == 0 else body.add_paragraph()
        p.text = bullet
        p.level = 0
        p.space_after = Pt(14)
        p.alignment = PP_ALIGN.LEFT
        if p.runs:
            run = p.runs[0]
            run.font.size = Pt(24)
            run.font.color.rgb = TEXT_COLOR

    add_speaker_notes(slide, notes)


def add_flow_slide(prs: Presentation, notes: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    style_background(slide)
    slide.shapes.title.text = "How The Agents Work Together"

    title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    title_run.font.size = Pt(30)
    title_run.font.bold = True
    title_run.font.color.rgb = TITLE_COLOR

    steps = [
        "Trip Request",
        "Safety Risk",
        "Telemetry",
        "Ops Monitoring",
        "Support",
        "Safe Outcome",
    ]

    x_start = 0.6
    y = 2.6
    box_w = 1.95
    box_h = 1.25
    gap = 0.23

    for i, step in enumerate(steps):
        x = x_start + i * (box_w + gap)
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(box_w),
            Inches(box_h),
        )
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_COLOR
        box.line.color.rgb = ACCENT_COLOR

        text_frame = box.text_frame
        text_frame.clear()
        p = text_frame.paragraphs[0]
        p.text = step
        p.alignment = PP_ALIGN.CENTER
        if p.runs:
            run = p.runs[0]
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = TITLE_COLOR

        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(x + box_w),
                Inches(y + 0.4),
                Inches(gap),
                Inches(0.45),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = MUTED_COLOR
            arrow.line.fill.background()

    subtitle = slide.shapes.add_textbox(
        Inches(0.9), Inches(4.45), Inches(11.8), Inches(1.15)
    )
    sub_tf = subtitle.text_frame
    sub_tf.clear()
    sp = sub_tf.paragraphs[0]
    sp.text = (
        "Safety is checked first, then monitored in real time, "
        "then resolved with structured support."
    )
    sp.alignment = PP_ALIGN.CENTER
    if sp.runs:
        srun = sp.runs[0]
        srun.font.size = Pt(20)
        srun.font.color.rgb = MUTED_COLOR

    add_speaker_notes(slide, notes)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Ride-Hailing Multi-Agent AI",
        "Demo-ready overview with speaker notes",
        (
            "Welcome. In this short presentation, I will show how our "
            "ride-hailing platform uses multiple AI agents to improve safety "
            "and operational trust."
        ),
    )

    add_bullet_slide(
        prs,
        "Business Goal",
        [
            "Improve passenger and driver safety",
            "Detect risky behavior in real time",
            "Resolve incidents faster and more fairly",
        ],
        (
            "Our business goal is to reduce risk without slowing operations. "
            "We focus on safety first, early risk detection during trips, and "
            "fair incident resolution for both riders and drivers."
        ),
    )

    add_bullet_slide(
        prs,
        "Four Specialized Agents",
        [
            "1) Safety risk assessment",
            "2) OBD telemetry and dangerous driving detection",
            "3) Live trip operations monitoring",
            "4) Customer support and incident resolution",
        ],
        (
            "Instead of one general model, we use four focused agents. "
            "Each has a clear responsibility, which makes decisions more "
            "reliable and easier to audit."
        ),
    )

    add_flow_slide(
        prs,
        (
            "Here is the full flow: trip request, pre-trip safety checks, "
            "live telemetry monitoring, in-trip operations oversight, and then "
            "support resolution when needed."
        ),
    )

    add_bullet_slide(
        prs,
        "Why This Architecture Works",
        [
            "Clear separation of responsibilities",
            "Better control, observability, and explainability",
            "Faster interventions and safer trips",
            "Easier to scale and improve over time",
        ],
        (
            "This architecture works because each agent is easier to tune, "
            "monitor, and improve. If something fails, we can isolate it "
            "quickly and recover without affecting the whole system."
        ),
    )

    add_bullet_slide(
        prs,
        "Expected Impact",
        [
            "Fewer safety incidents",
            "Higher rider trust and retention",
            "Reduced support workload",
            "Stronger operational decisions from live data",
        ],
        (
            "The expected result is safer trips, better customer trust, lower "
            "support pressure, and stronger decisions because operations are "
            "based on live evidence, not only historical reports."
        ),
    )

    prs.save(str(OUTPUT_FILE))
    print(f"Presentation generated: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()

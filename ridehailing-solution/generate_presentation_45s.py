from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_FILE = BASE_DIR / "ridehailing-multi-agent-presentation-45s.pptx"


TITLE_COLOR = RGBColor(20, 45, 75)
ACCENT_COLOR = RGBColor(0, 130, 165)
TEXT_COLOR = RGBColor(35, 35, 35)
BG_COLOR = RGBColor(245, 249, 253)
CARD_COLOR = RGBColor(255, 255, 255)


def style_bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        Inches(13.333),
        Inches(0.55),
    )
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(12, 32, 56)
    band.line.fill.background()


def add_notes(slide, text: str) -> None:
    notes = slide.notes_slide.notes_text_frame
    notes.clear()
    notes.text = text


def add_title(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    style_bg(slide)

    slide.shapes.title.text = "Ride-Hailing Multi-Agent AI"
    slide.placeholders[1].text = "45-second pitch"

    t_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    t_run.font.size = Pt(44)
    t_run.font.bold = True
    t_run.font.color.rgb = TITLE_COLOR

    s_run = slide.placeholders[1].text_frame.paragraphs[0].runs[0]
    s_run.font.size = Pt(22)
    s_run.font.color.rgb = ACCENT_COLOR

    add_notes(
        slide,
        "This project uses multiple AI agents to make ride-hailing safer, faster, and more reliable.",
    )


def add_bullets(prs: Presentation, title: str, bullets: list[str], notes: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    style_bg(slide)

    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.6),
        Inches(1.15),
        Inches(12.15),
        Inches(5.95),
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = CARD_COLOR
    panel.line.color.rgb = RGBColor(220, 228, 238)

    slide.shapes.title.text = title
    tr = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    tr.font.size = Pt(34)
    tr.font.bold = True
    tr.font.color.rgb = TITLE_COLOR

    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(16)
        if p.runs:
            run = p.runs[0]
            run.font.size = Pt(24)
            run.font.color.rgb = TEXT_COLOR

    add_notes(slide, notes)


def add_flow(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    style_bg(slide)

    slide.shapes.title.text = "How It Works"
    tr = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    tr.font.size = Pt(34)
    tr.font.bold = True
    tr.font.color.rgb = TITLE_COLOR

    steps = ["Request", "Risk", "Telemetry", "Support"]
    x0 = 1.0
    y = 2.3
    w = 2.6
    h = 1.5
    gap = 0.5

    for i, step in enumerate(steps):
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x0 + i * (w + gap)),
            Inches(y),
            Inches(w),
            Inches(h),
        )
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_COLOR
        box.line.color.rgb = ACCENT_COLOR

        tf = box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = step
        p.alignment = PP_ALIGN.CENTER
        if p.runs:
            run = p.runs[0]
            run.font.size = Pt(24)
            run.font.bold = True
            run.font.color.rgb = TITLE_COLOR

        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.CHEVRON,
                Inches(x0 + w + i * (w + gap)),
                Inches(y + 0.45),
                Inches(gap - 0.1),
                Inches(0.6),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = ACCENT_COLOR
            arrow.line.fill.background()

    summary = slide.shapes.add_textbox(Inches(1.2), Inches(4.5), Inches(10.8), Inches(1.0))
    s_tf = summary.text_frame
    s_tf.clear()
    sp = s_tf.paragraphs[0]
    sp.text = "Early detection + fast escalation = safer trips"
    sp.alignment = PP_ALIGN.CENTER
    if sp.runs:
        run = sp.runs[0]
        run.font.size = Pt(22)
        run.font.color.rgb = ACCENT_COLOR

    add_notes(
        slide,
        "Flow is simple: request, risk check, live telemetry watch, and support action when needed.",
    )


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title(prs)

    add_bullets(
        prs,
        "The Problem",
        [
            "Safety issues can appear suddenly",
            "One general AI is too broad",
            "Operations need real-time decisions",
        ],
        "The challenge is safety at scale. We need quick and controlled decisions, not one generic model.",
    )

    add_flow(prs)

    add_bullets(
        prs,
        "Impact",
        [
            "Fewer incidents",
            "Faster resolution",
            "Higher rider trust",
        ],
        "Result: safer rides, faster support, and stronger trust for riders and drivers.",
    )

    prs.save(str(OUTPUT_FILE))
    print(f"Presentation generated: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()

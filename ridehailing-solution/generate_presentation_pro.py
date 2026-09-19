from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_FILE = BASE_DIR / "ridehailing-multi-agent-presentation-pro.pptx"


NAVY = RGBColor(11, 30, 54)
SKY = RGBColor(0, 140, 170)
TEAL = RGBColor(0, 168, 150)
AMBER = RGBColor(245, 167, 66)
INK = RGBColor(30, 35, 46)
MUTED = RGBColor(83, 98, 122)
BG = RGBColor(240, 246, 252)
WHITE = RGBColor(255, 255, 255)


class DeckBuilder:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)

    def apply_background(self, slide, accent: RGBColor = SKY) -> None:
        bg_fill = slide.background.fill
        bg_fill.solid()
        bg_fill.fore_color.rgb = BG

        top_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(0),
            Inches(13.333),
            Inches(0.58),
        )
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = NAVY
        top_bar.line.fill.background()

        # Decorative circles create depth while keeping slide content readable.
        c1 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.8), Inches(-0.7), Inches(3.2), Inches(3.2))
        c1.fill.solid()
        c1.fill.fore_color.rgb = accent
        c1.line.fill.background()

        c2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-0.6), Inches(5.9), Inches(2.6), Inches(2.6))
        c2.fill.solid()
        c2.fill.fore_color.rgb = RGBColor(210, 228, 247)
        c2.line.fill.background()

    def set_title(self, slide, title: str, subtitle: str = "") -> None:
        title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.78), Inches(12.0), Inches(0.9))
        tf = title_box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.LEFT
        if p.runs:
            run = p.runs[0]
            run.font.size = Pt(36)
            run.font.bold = True
            run.font.color.rgb = NAVY

        if subtitle:
            sbox = slide.shapes.add_textbox(Inches(0.72), Inches(1.55), Inches(11.8), Inches(0.55))
            stf = sbox.text_frame
            stf.clear()
            sp = stf.paragraphs[0]
            sp.text = subtitle
            if sp.runs:
                srun = sp.runs[0]
                srun.font.size = Pt(18)
                srun.font.color.rgb = MUTED

    def add_notes(self, slide, notes: str) -> None:
        nf = slide.notes_slide.notes_text_frame
        nf.clear()
        nf.text = notes

    def add_title_slide(self) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.apply_background(slide, accent=TEAL)

        hero = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.9),
            Inches(1.45),
            Inches(11.6),
            Inches(4.95),
        )
        hero.fill.solid()
        hero.fill.fore_color.rgb = WHITE
        hero.line.color.rgb = RGBColor(215, 227, 240)

        title = slide.shapes.add_textbox(Inches(1.3), Inches(2.1), Inches(10.6), Inches(1.6))
        tf = title.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = "Multi-Agent AI for Ride-Hailing Safety"
        if p.runs:
            run = p.runs[0]
            run.font.size = Pt(44)
            run.font.bold = True
            run.font.color.rgb = NAVY

        p2 = tf.add_paragraph()
        p2.text = "Production-ready architecture based on solution.md"
        if p2.runs:
            r2 = p2.runs[0]
            r2.font.size = Pt(23)
            r2.font.color.rgb = SKY

        chips = [
            "Safety & Risk", "OBD Telemetry", "Trip Operations", "Incident Resolution"
        ]
        x = 1.35
        for chip in chips:
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x),
                Inches(4.65),
                Inches(2.58),
                Inches(0.68),
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = BG
            shape.line.color.rgb = RGBColor(208, 221, 235)
            ctf = shape.text_frame
            ctf.clear()
            cp = ctf.paragraphs[0]
            cp.text = chip
            cp.alignment = PP_ALIGN.CENTER
            if cp.runs:
                cr = cp.runs[0]
                cr.font.size = Pt(15)
                cr.font.bold = True
                cr.font.color.rgb = INK
            x += 2.75

        self.add_notes(
            slide,
            "This presentation summarizes the ride-hailing solution using four specialized agents for safety, telemetry, operations, and support.",
        )

    def add_problem_slide(self) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.apply_background(slide, accent=SKY)
        self.set_title(
            slide,
            "Business Problem",
            "Ride-hailing safety and trust need real-time, policy-aligned decisions.",
        )

        left = self._add_card(slide, 0.8, 2.0, 6.0, 4.9, "Key Challenges")
        right = self._add_card(slide, 6.6, 2.0, 5.9, 4.9, "Why Single Agent Fails")

        self._add_list(
            left,
            [
                "Safety incidents can happen suddenly during trips",
                "GPS-only signals miss dangerous driving patterns",
                "Support teams need fast and traceable escalation",
                "Fraud, disputes, and policy consistency must be managed together",
            ],
            18,
        )
        self._add_list(
            right,
            [
                "Too broad to optimize for high-risk decisions",
                "Harder to audit and debug errors",
                "Weak separation between safety and support logic",
                "Lower reliability under live operational pressure",
            ],
            18,
        )

        self.add_notes(
            slide,
            "The core issue is that safety, operations, and support are connected. One general model is hard to control, so we split responsibilities.",
        )

    def add_agents_slide(self) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.apply_background(slide, accent=TEAL)
        self.set_title(
            slide,
            "Four Specialized Agents",
            "Each agent has clear tools, decisions, and escalation behavior.",
        )

        cards = [
            (0.8, 2.0, "1. Safety Risk Agent", "Profiles, route risk, pre-checks"),
            (6.8, 2.0, "2. Telemetry Agent", "OBD speed, RPM, braking anomalies"),
            (0.8, 4.35, "3. Operations Agent", "Detours, idle periods, trip monitoring"),
            (6.8, 4.35, "4. Support Agent", "Complaints, policy checks, resolution"),
        ]

        for x, y, title, subtitle in cards:
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x),
                Inches(y),
                Inches(5.7),
                Inches(2.0),
            )
            card.fill.solid()
            card.fill.fore_color.rgb = WHITE
            card.line.color.rgb = RGBColor(211, 224, 238)

            t = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.25), Inches(5.2), Inches(0.8))
            tf = t.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = title
            if p.runs:
                r = p.runs[0]
                r.font.size = Pt(22)
                r.font.bold = True
                r.font.color.rgb = NAVY

            s = tf.add_paragraph()
            s.text = subtitle
            if s.runs:
                sr = s.runs[0]
                sr.font.size = Pt(16)
                sr.font.color.rgb = MUTED

        self.add_notes(
            slide,
            "This separation gives us better control and makes every decision point measurable and explainable.",
        )

    def add_main_flow_slide(self) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.apply_background(slide, accent=SKY)
        self.set_title(slide, "End-to-End Flow Diagram", "Derived directly from section 5 in solution.md")

        steps = [
            "Trip Request",
            "Safety Risk\nAssessment",
            "OBD Telemetry\nMonitoring",
            "In-Trip Ops\nMonitoring",
            "Support &\nResolution",
            "Business\nOutcome",
        ]

        x = 0.55
        y = 2.55
        w = 2.0
        h = 1.45
        gap = 0.2

        for i, step in enumerate(steps):
            box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x + i * (w + gap)),
                Inches(y),
                Inches(w),
                Inches(h),
            )
            box.fill.solid()
            box.fill.fore_color.rgb = WHITE
            box.line.color.rgb = SKY if i % 2 == 0 else TEAL
            tf = box.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = step
            p.alignment = PP_ALIGN.CENTER
            if p.runs:
                r = p.runs[0]
                r.font.size = Pt(16)
                r.font.bold = True
                r.font.color.rgb = NAVY

            if i < len(steps) - 1:
                arr = slide.shapes.add_shape(
                    MSO_SHAPE.RIGHT_ARROW,
                    Inches(x + w + i * (w + gap)),
                    Inches(y + 0.47),
                    Inches(gap),
                    Inches(0.52),
                )
                arr.fill.solid()
                arr.fill.fore_color.rgb = RGBColor(132, 153, 179)
                arr.line.fill.background()

        caption = slide.shapes.add_textbox(Inches(0.85), Inches(4.35), Inches(11.8), Inches(1.6))
        ctf = caption.text_frame
        ctf.clear()
        cp = ctf.paragraphs[0]
        cp.text = "Safety is checked before acceptance, monitored in real time, then resolved through policy-grounded support."
        cp.alignment = PP_ALIGN.CENTER
        if cp.runs:
            cr = cp.runs[0]
            cr.font.size = Pt(19)
            cr.font.color.rgb = MUTED

        self.add_notes(
            slide,
            "This is the main production flow: pre-trip risk scoring, live telemetry checks, in-trip anomaly handling, and structured support outcomes.",
        )

    def add_escalation_flow_slide(self) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.apply_background(slide, accent=AMBER)
        self.set_title(slide, "Escalation Scenario Flow", "Based on the suspicious trip example in solution.md")

        swim = [
            ("Passenger", RGBColor(233, 245, 255)),
            ("Safety Agent", RGBColor(228, 252, 247)),
            ("Telemetry Agent", RGBColor(255, 245, 227)),
            ("Operations", RGBColor(245, 238, 255)),
            ("Support", RGBColor(255, 236, 240)),
        ]

        y = 2.0
        lane_h = 0.86
        for label, color in swim:
            lane = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(0.75),
                Inches(y),
                Inches(12.0),
                Inches(lane_h),
            )
            lane.fill.solid()
            lane.fill.fore_color.rgb = color
            lane.line.color.rgb = RGBColor(209, 219, 231)

            lt = slide.shapes.add_textbox(Inches(0.95), Inches(y + 0.2), Inches(1.8), Inches(0.4))
            ltf = lt.text_frame
            ltf.clear()
            lp = ltf.paragraphs[0]
            lp.text = label
            if lp.runs:
                lr = lp.runs[0]
                lr.font.size = Pt(14)
                lr.font.bold = True
                lr.font.color.rgb = NAVY
            y += 0.92

        events = [
            (3.0, 2.0, "Late-night\nrequest"),
            (4.7, 2.92, "High route\nrisk score"),
            (6.45, 3.84, "Harsh braking\n+ speeding"),
            (8.2, 4.76, "Manual\nintervention"),
            (9.95, 5.68, "Case review\n+ resolution"),
        ]

        for i, (x, y_ev, text) in enumerate(events):
            event = slide.shapes.add_shape(
                MSO_SHAPE.CHEVRON,
                Inches(x),
                Inches(y_ev + 0.07),
                Inches(1.65),
                Inches(0.72),
            )
            event.fill.solid()
            event.fill.fore_color.rgb = WHITE
            event.line.color.rgb = AMBER
            etf = event.text_frame
            etf.clear()
            ep = etf.paragraphs[0]
            ep.text = text
            ep.alignment = PP_ALIGN.CENTER
            if ep.runs:
                er = ep.runs[0]
                er.font.size = Pt(12)
                er.font.bold = True
                er.font.color.rgb = INK

            if i < len(events) - 1:
                arrow = slide.shapes.add_shape(
                    MSO_SHAPE.DOWN_ARROW,
                    Inches(x + 1.3),
                    Inches(y_ev + 0.82),
                    Inches(0.24),
                    Inches(0.24),
                )
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = RGBColor(165, 178, 196)
                arrow.line.fill.background()

        self.add_notes(
            slide,
            "In this scenario, telemetry confirms risky behavior, operations intervenes, and support finalizes a policy-aligned case outcome.",
        )

    def add_readiness_slide(self) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.apply_background(slide, accent=TEAL)
        self.set_title(slide, "Production Readiness", "Observability, evaluation, and safety controls")

        l = self._add_card(slide, 0.8, 2.0, 4.0, 4.9, "Observability")
        m = self._add_card(slide, 4.95, 2.0, 4.0, 4.9, "Evaluation")
        r = self._add_card(slide, 9.1, 2.0, 3.95, 4.9, "Controls")

        self._add_list(l, ["Trace every agent call", "Monitor latency and cost", "Track anomalies and escalations"], 16)
        self._add_list(m, ["Use scenario-rich test datasets", "Measure safety and groundedness", "Run checks before every release"], 16)
        self._add_list(r, ["Policy-grounded actions", "Human review for critical cases", "Audit logs for key decisions"], 16)

        self.add_notes(
            slide,
            "This is what turns a demo into a deployable system: monitoring, evaluation discipline, and operational safeguards.",
        )

    def add_closing_slide(self) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.apply_background(slide, accent=SKY)

        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(1.2),
            Inches(1.7),
            Inches(10.9),
            Inches(4.5),
        )
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = RGBColor(216, 228, 241)

        t = slide.shapes.add_textbox(Inches(1.55), Inches(2.2), Inches(10.1), Inches(2.8))
        tf = t.text_frame
        tf.clear()

        p = tf.paragraphs[0]
        p.text = "Outcome"
        if p.runs:
            r = p.runs[0]
            r.font.size = Pt(30)
            r.font.bold = True
            r.font.color.rgb = NAVY

        p2 = tf.add_paragraph()
        p2.text = "Safer trips, faster intervention, and trusted support decisions"
        if p2.runs:
            r2 = p2.runs[0]
            r2.font.size = Pt(24)
            r2.font.color.rgb = SKY

        p3 = tf.add_paragraph()
        p3.text = "Multi-agent architecture makes the platform explainable, scalable, and production-ready."
        if p3.runs:
            r3 = p3.runs[0]
            r3.font.size = Pt(19)
            r3.font.color.rgb = MUTED

        self.add_notes(
            slide,
            "In summary, the solution combines AI specialization with operational safety and production controls to improve both trust and performance.",
        )

    def _add_card(self, slide, x: float, y: float, w: float, h: float, heading: str):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(w),
            Inches(h),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = RGBColor(214, 226, 239)

        heading_box = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.18), Inches(w - 0.5), Inches(0.55))
        tf = heading_box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = heading
        if p.runs:
            run = p.runs[0]
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = NAVY

        body = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.8), Inches(w - 0.45), Inches(h - 1.0))
        return body

    def _add_list(self, textbox, lines: list[str], size: int) -> None:
        tf = textbox.text_frame
        tf.clear()
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = "- " + line
            p.space_after = Pt(10)
            if p.runs:
                run = p.runs[0]
                run.font.size = Pt(size)
                run.font.color.rgb = INK

    def build(self) -> None:
        self.add_title_slide()
        self.add_problem_slide()
        self.add_agents_slide()
        self.add_main_flow_slide()
        self.add_escalation_flow_slide()
        self.add_readiness_slide()
        self.add_closing_slide()
        self.prs.save(str(OUTPUT_FILE))


def main() -> None:
    builder = DeckBuilder()
    builder.build()
    print(f"Presentation generated: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
